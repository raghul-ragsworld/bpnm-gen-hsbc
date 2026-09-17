import argparse
from collections import Counter, defaultdict
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re

import openpyxl
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parent
DEFAULT_SOURCE = ROOT / "Synthetic_Hackathon_Data_10.xlsx"
PRIMARY_KEYS = {
    "Process": "Process ID", "Process Step": "Process Step ID",
    "Process Step Detail": "PSD ID", "Process Step Data": "PSDATA ID",
    "Application": "Application ID", "Attribute Catalogue": "Attribute ID",
    "Process ID Mapping": "Process ID",
}
TEAM_ALIASES = {
    "Requestor (Client Service)": "Requestor", "Onboarding Ops (Maker)": "Onboarding Ops",
    "Maker": "Onboarding Ops", "Process Ops (DQ & Controls)": "Process Ops",
    "Process Ops (Checker)": "Process Ops", "Checker": "Process Ops",
    "Business Approver": "Approver", "Integration Layer": "Integration Team",
    "Integration": "Integration Team", "Data Platform": "Data Platform Team",
    "Platform": "Data Platform Team", "Service Management (SNOW)": "Service Management Team",
    "Service Management (ServiceNow Team)": "Service Management Team",
    "SNOW": "Service Management Team", "ServiceNow": "Service Management Team",
    "Data Steward (Reference Data)": "Steward", "Data Steward": "Steward",
    "Ops": "Operations Team (unresolved)",
    "Downstream Systems": "Downstream Consumer Team (proposed)",
    "DQ Dashboard": "Reporting Team (proposed)",
}
APPLICATION_CAPABILITIES = {
    "APP-001": "Request & Workflow", "APP-002": "Request & Workflow",
    "APP-003": "Integration & Delivery", "APP-004": "Integration & Delivery",
    "APP-005": "Integration & Delivery", "APP-006": "Data Platform",
    "APP-007": "Data Quality", "APP-008": "Data Platform",
    "APP-009": "Incident Management", "APP-010": "External Data / Consumers",
    "APP-011": "External Data / Consumers", "APP-012": "External Data / Consumers",
    "APP-013": "Reporting",
}
STEP_REFERENCE = re.compile(r"\bAOB-PR-\d+-V\d+-S\d{3}\b")
LEGACY_REFERENCE = re.compile(r"\bAOB-\d{3}\b")


def tables_from(workbook):
    tables = {}
    for sheet in workbook:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        tables[sheet.title] = [
            {**dict(zip(headers, values)), "_row": number}
            for number, values in enumerate(rows, 2) if any(value is not None for value in values)
        ]
    return tables


def parse_raci(value):
    assignments = {}
    for part in str(value or "").split(","):
        role, separator, team = part.strip().partition("=")
        if not separator or role not in {"R", "A", "C", "I"} or not team.strip() or role in assignments:
            raise ValueError(f"Invalid RACI: {value!r}")
        assignments[role] = team.strip()
    if set(assignments) != {"R", "A", "C", "I"}:
        raise ValueError(f"Incomplete RACI: {value!r}")
    return assignments


def audit(tables):
    errors = []
    indexes = {}
    for sheet, key in PRIMARY_KEYS.items():
        records = tables[sheet]
        counts = Counter(record[key] for record in records)
        for value, count in counts.items():
            if not value or count != 1:
                errors.append({"sheet": sheet, "key": key, "value": value, "count": count})
        indexes[sheet] = {record[key]: record for record in records}
    references = {"Process ID": "Process", "Process Step ID": "Process Step",
                  "PSD ID": "Process Step Detail", "PSDATA ID": "Process Step Data",
                  "Application ID": "Application", "Attribute ID": "Attribute Catalogue"}
    for sheet, records in tables.items():
        if sheet in {"Source Process", "Generation Notes"}:
            continue
        for record in records:
            for key, target in references.items():
                if key in record and record[key] not in indexes[target]:
                    errors.append({"sheet": sheet, "row": record["_row"], "missing": key, "value": record[key]})
            step = indexes["Process Step"].get(record.get("Process Step ID"))
            if step:
                for key in ("Process ID", "Process Step"):
                    if key in record and record[key] != step[key]:
                        errors.append({"sheet": sheet, "row": record["_row"], "mismatch": key})
            process = indexes["Process"].get(record.get("Process ID"))
            if process and "L4 Process Name" in record and record["L4 Process Name"] != process["L4 Process Name"]:
                errors.append({"sheet": sheet, "row": record["_row"], "mismatch": "L4 Process Name"})
            for key, target, fields in (
                ("PSD ID", "Process Step Detail", ("Process Step ID",)),
                ("PSDATA ID", "Process Step Data", ("Process Step ID", "BDE Name", "Source/Target")),
                ("Application ID", "Application", ("Application Name",)),
                ("Attribute ID", "Attribute Catalogue", ("Attribute Name",)),
            ):
                parent = indexes[target].get(record.get(key))
                if parent:
                    for field in fields:
                        if field in record and record[field] != parent[field]:
                            errors.append({"sheet": sheet, "row": record["_row"], "mismatch": field})
            if "RACI" in record:
                try:
                    parse_raci(record["RACI"])
                except ValueError as error:
                    errors.append({"sheet": sheet, "row": record["_row"], "error": str(error)})
            if "Source/Target" in record and record["Source/Target"] not in {"Source", "Target"}:
                errors.append({"sheet": sheet, "row": record["_row"], "error": "Invalid Source/Target"})
            context_process = record.get("Process ID") or (step or {}).get("Process ID")
            for field, value in record.items():
                if not isinstance(value, str) or field.startswith("Source "):
                    continue
                for reference in STEP_REFERENCE.findall(value):
                    target_step = indexes["Process Step"].get(reference)
                    if not target_step or (context_process and target_step["Process ID"] != context_process):
                        errors.append({"sheet": sheet, "row": record["_row"], "error": "Invalid/cross-process step reference", "value": reference})
                if LEGACY_REFERENCE.search(value):
                    errors.append({"sheet": sheet, "row": record["_row"], "error": "Unqualified legacy step reference"})
    for sheet, fields in (
        ("Step Application Mapping", ("PSD ID", "Application ID")),
        ("Step Attribute Mapping", ("PSDATA ID", "Attribute ID")),
        ("Process Step", ("Process ID", "Reorder")),
    ):
        for key, count in Counter(tuple(row[field] for field in fields) for row in tables[sheet]).items():
            if count != 1:
                errors.append({"sheet": sheet, "error": "Duplicate composite key", "key": key})
    app_details = {row["PSD ID"] for row in tables["Process Step Detail"] if row["Detail Type"] == "Application"}
    mapped_details = {row["PSD ID"] for row in tables["Step Application Mapping"]}
    if app_details != mapped_details:
        errors.append({"error": "Application detail mapping coverage mismatch", "ids": sorted(app_details ^ mapped_details)})
    data_ids = set(indexes["Process Step Data"])
    mapped_data = {row["PSDATA ID"] for row in tables["Step Attribute Mapping"]}
    if data_ids != mapped_data:
        errors.append({"error": "Attribute mapping coverage mismatch", "ids": sorted(data_ids ^ mapped_data)})
    source_rows = {row["_row"]: row for row in tables["Source Process"]}
    mapped_processes = {row["Process ID"] for row in tables["Process ID Mapping"]}
    if mapped_processes != set(indexes["Process"]):
        errors.append({"error": "Process provenance mapping coverage mismatch"})
    processes_with_steps = {row["Process ID"] for row in tables["Process Step"]}
    if processes_with_steps != set(indexes["Process"]):
        errors.append({"error": "Process/step coverage mismatch"})
    for mapping in tables["Process ID Mapping"]:
        source = source_rows.get(mapping["Source Excel Row"])
        if not source or (source["Process ID"], source["Process Name"]) != (mapping["Source Process ID"], mapping["Source Process Name"]):
            errors.append({"sheet": "Process ID Mapping", "row": mapping["_row"], "error": "Source provenance mismatch"})
    details = defaultdict(list)
    for record in tables["Process Step Detail"]:
        details[record["Process Step ID"]].append(record)
    missing_apps = []
    missing_data = []
    data_steps = {record["Process Step ID"] for record in tables["Process Step Data"]}
    for record in tables["Process Step"]:
        identifier = record["Process Step ID"]
        if not details[identifier]:
            errors.append({"sheet": "Process Step", "row": record["_row"], "error": "No step details"})
        if not isinstance(record["Reorder"], (int, float)) or record["Reorder"] <= 0:
            errors.append({"sheet": "Process Step", "row": record["_row"], "error": "Invalid display order"})
        if not any(detail["Detail Type"] == "Application" for detail in details[identifier]):
            missing_apps.append(identifier)
        if identifier not in data_steps:
            missing_data.append(identifier)
    return {"counts": {sheet: len(records) for sheet, records in tables.items()},
            "structural_errors": errors, "missing_application_steps": missing_apps,
            "missing_data_steps": missing_data}


def normalize_assignments(owner, raw_raci):
    source = parse_raci(raw_raci)
    team = TEAM_ALIASES.get(owner, owner)
    assignments = {role: TEAM_ALIASES.get(value, value) for role, value in source.items()}
    if assignments["R"] != team and source["R"] != "Ops":
        raise ValueError(f"Owner/R conflict requires review: {owner!r}, {raw_raci!r}")
    assignments["R"] = team
    duties = {role: value for role, value in source.items() if value in {"Maker", "Checker"}}
    return team, ", ".join(f"{role}={assignments[role]}" for role in ("R", "A", "C", "I")), duties


def normalize(tables):
    result = deepcopy(tables)
    changes = []
    for row in result["Process Step Detail"]:
        owner, raw_raci = row["Process Step Owner"], row["RACI"]
        team, raci, duties = normalize_assignments(owner, raw_raci)
        row.update({"Source Process Step Owner": owner, "Source RACI": raw_raci,
                    "Source Detail Type": row["Detail Type"], "Source Detail Value": row["Detail Value"],
                    "Control Duty": duties.get("R", ""), "RACI Duties": json.dumps(duties, sort_keys=True)})
        replacements = {"Process Step Owner": team, "RACI": raci}
        if row["Detail Type"] == "Gateway Rule (Inclusive)":
            if not re.fullmatch(r"If Tax Review Basis = TRUE then route to (AOB-PR-\d+-V\d+-S\d{3}); otherwise continue to (AOB-PR-\d+-V\d+-S\d{3})", row["Detail Value"]):
                raise ValueError("Inclusive gateway requires manual review: " + row["PSD ID"])
            replacements["Detail Type"] = "Gateway Rule (XOR)"
        for field, value in replacements.items():
            if row[field] != value:
                changes.append({"Sheet": "Process Step Detail", "Row": row["_row"], "ID": row["PSD ID"],
                                "Column": field, "Before": row[field], "After": value,
                                "Reason": "Mutually exclusive if/otherwise branches" if field == "Detail Type" else "Human-team normalization; source and duties preserved"})
                row[field] = value
    return result, changes


def build_metadata(tables):
    details = defaultdict(list)
    applications = defaultdict(list)
    data = defaultdict(list)
    for row in tables["Process Step Detail"]:
        details[row["Process Step ID"]].append(row)
    for row in tables["Step Application Mapping"]:
        applications[row["Process Step ID"]].append(row)
    for row in tables["Process Step Data"]:
        data[row["Process Step ID"]].append(row)
    steps, raci_rows, controls, issues = [], [], [], []

    def issue(code, process_id, step_id, message):
        issues.append({"Code": code, "Process ID": process_id, "Process Step ID": step_id,
                       "Status": "Needs confirmation", "Description": message})

    for process in tables["Process"]:
        issue("SYNTHETIC_TEMPLATE", process["Process ID"], "", "Shared onboarding template is synthetic, not a verified domain-specific workflow.")
        issue("NO_EXPLICIT_FLOW_TABLE", process["Process ID"], "", "Reorder is display order, not sequence flow. Complete start/end, branch, merge, retry and timer semantics must be modeled before BPMN graph validation.")
    for step in tables["Process Step"]:
        identifier, process_id = step["Process Step ID"], step["Process ID"]
        records = details[identifier]
        teams = sorted({row["Process Step Owner"] for row in records})
        if len(teams) != 1:
            raise ValueError(f"Multiple responsible teams: {identifier}: {teams}")
        types = {row["Detail Type"] for row in records}
        kind = "task"
        if "Gateway Rule (XOR)" in types:
            kind = "exclusiveGateway"
        elif types & {"Gateway (AND)", "Gateway (AND Join)"}:
            kind = "parallelGateway"
        elif "Event-based Gateway" in types:
            kind = "eventBasedGateway"
        app_ids = sorted({row["Application ID"] for row in applications[identifier]})
        app_values = [row["Detail Value"] for row in records if row["Detail Type"] == "Application"]
        if not app_values:
            issue("MISSING_APPLICATION", process_id, identifier, "No explicit application detail. Hosting/execution mode is unknown; no assignment fabricated.")
        if not data[identifier]:
            issue("NO_DATA_MAPPING", process_id, identifier, "No source BDE mapping; this does not remove the step or imply an invalid BPMN task.")
        seen = set()
        for detail in records:
            variant = (detail["Source Process Step Owner"], detail["Source RACI"])
            if variant not in seen:
                seen.add(variant)
                assignments = parse_raci(detail["RACI"])
                raci_rows.append({"Process ID": process_id, "Process Step ID": identifier,
                                  "Variant": len(seen), "Responsible Team": teams[0], **assignments,
                                  "Control Duty": detail["Control Duty"], "RACI Duties": detail["RACI Duties"],
                                  "Source Owner": variant[0], "Source RACI": variant[1]})
                if any("(unresolved)" in value or "(proposed)" in value for value in assignments.values()):
                    issue("RACI_CONFIRMATION", process_id, identifier, "Unresolved Operations team or proposed consumer/reporting team in RACI; original roles retained.")
            text = detail["Detail Value"]
            references = list(dict.fromkeys(STEP_REFERENCE.findall(text)))
            if references or any(word in detail["Detail Type"] for word in ("Gateway", "Timer", "Rule", "Control", "Exception", "Transaction")):
                controls.append({"Process ID": process_id, "Process Step ID": identifier, "PSD ID": detail["PSD ID"],
                                 "Control Type": detail["Detail Type"], "Rule Text": text,
                                 "Referenced Step IDs": references,
                                 "Interpretation": "Source references only, not executable sequence flows"})
        all_text = " ".join(str(row["Detail Value"]) for row in records)
        if "Event-based Gateway" in types:
            issue("APPROVAL_WAIT_MODEL", process_id, identifier, "Model approval/rejection/withdrawal as competing catches; attach the 24h boundary timer to an enclosing wait activity/subprocess, not a gateway. Confirm timeout-review outcome for this variant.")
        if "Loop Rule" in types:
            issue("REMEDIATION_OUTCOMES", process_id, identifier, "Separate urgent 4h/standard 24h timer breach from 3-attempt exhaustion. Confirm terminal parking/escalation paths for this variant.")
        if any("Checker" in row["Control Duty"] for row in records):
            issue("CHECKER_REJECTION", process_id, identifier, "Approval is required; rejected-checker destination is unspecified. Maker and Checker must be different individuals for the same record; staffing evidence is not provided.")
        if "Tax Review Basis is derived" in all_text:
            issue("TAX_RULE_VOCABULARY", process_id, identifier, "TAXABLE is used as Process Status; allowed status vocabulary and tax-review-required type list are missing.")
        if "Responsible Groups" in all_text:
            issue("ATTRIBUTE_ALIAS", process_id, identifier, "Responsible Groups versus Process Teams remains an unconfirmed alias.")
        if "Azure Monitor Alert / Incident Record" in all_text:
            issue("INCIDENT_APPLICATION", process_id, identifier, "Source mixes Azure Monitor incident wording, ServiceNow data mappings and Incident Management Tool inventory. Do not silently equate these applications.")
        if "Enriched Record" in str(step["Source BDEs"]):
            issue("PARALLEL_DATA_DEPENDENCY", process_id, identifier, "Duplicate detection consumes Enriched Record while the source AND split runs it alongside enrichment. Execution order needs confirmation.")
        duties = sorted({row["Control Duty"] for row in records if row["Control Duty"]})
        steps.append({"Process ID": process_id, "Process Step ID": identifier, "Process Step": step["Process Step"],
                      "Display Order": step["Reorder"], "People Lane": teams[0], "Control Duties": duties,
                      "BPMN Element Candidate": kind, "Execution Mode": "Unspecified; not inferred from application presence",
                      "Application IDs": app_ids, "Source Application Details": app_values,
                      "Technology Capabilities": sorted({APPLICATION_CAPABILITIES[app_id] for app_id in app_ids}),
                      "Application Status": "Source-linked" if app_values else "Missing source assignment",
                      "Application Semantics": "Preserve source alternatives/composites/qualifiers; multiple links do not imply parallel execution",
                      "Data Mapping Count": len(data[identifier]), "RACI Variants": len(seen),
                      "Source PSD IDs": [row["PSD ID"] for row in records],
                      "Source PSDATA IDs": [row["PSDATA ID"] for row in data[identifier]]})
    unique_issues = list({(row["Code"], row["Process ID"], row["Process Step ID"]): row for row in issues}.values())
    return {"steps": steps, "raci": raci_rows, "controls": controls, "issues": unique_issues}


def verify_normalization(original, normalized, changes, metadata):
    errors = audit(normalized)["structural_errors"]
    logged = {(row["Sheet"], row["Row"], row["Column"]): row for row in changes}
    for sheet, records in original.items():
        if len(records) != len(normalized[sheet]):
            errors.append({"error": "Source row count changed", "sheet": sheet})
            continue
        for before, after in zip(records, normalized[sheet]):
            for field, value in before.items():
                if after[field] != value:
                    change = logged.get((sheet, before["_row"], field))
                    if not change or change["Before"] != value or change["After"] != after[field]:
                        errors.append({"error": "Unlogged source edit", "sheet": sheet, "row": before["_row"], "field": field})
            if sheet == "Process Step Detail":
                for field in ("Process Step Owner", "RACI", "Detail Type", "Detail Value"):
                    if after.get("Source " + field) != before[field]:
                        errors.append({"error": "Source provenance lost", "row": before["_row"], "field": field})
    for row in normalized["Process Step Detail"]:
        if parse_raci(row["RACI"])["R"] != row["Process Step Owner"]:
            errors.append({"error": "Responsible team mismatch", "id": row["PSD ID"]})
        source = parse_raci(row["Source RACI"])
        if source["R"] in {"Maker", "Checker"} and row["Control Duty"] != source["R"]:
            errors.append({"error": "Control duty lost", "id": row["PSD ID"]})
    if {row["Process Step ID"] for row in metadata["steps"]} != {row["Process Step ID"] for row in original["Process Step"]}:
        errors.append({"error": "Metadata step coverage mismatch"})
    for row in metadata["steps"]:
        if row["People Lane"] in {"Integration Layer", "Data Platform", "SNOW", "ServiceNow"}:
            errors.append({"error": "System in human lane", "id": row["Process Step ID"]})
    return errors


def excel_value(value):
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=True, sort_keys=True)
    return value


def write_sheet(workbook, name, records):
    sheet = workbook.create_sheet(name)
    if not records:
        sheet.append(["No records"])
        return
    headers = [key for key in records[0] if key != "_row"]
    sheet.append(headers)
    for row in records:
        sheet.append([excel_value(row.get(header)) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="285840")
        sheet.column_dimensions[cell.column_letter].width = min(48, max(18, len(str(cell.value)) + 3))


def run(source, output_dir):
    source = source.resolve()
    output_dir = output_dir.resolve()
    destination = output_dir / f"{source.stem}_normalized.xlsx"
    if destination == source:
        raise ValueError("Output cannot replace source workbook")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(source, data_only=False)
    try:
        original = tables_from(workbook)
        before = audit(original)
        if before["structural_errors"]:
            raise ValueError(json.dumps(before["structural_errors"], ensure_ascii=True))
        normalized, changes = normalize(original)
        metadata = build_metadata(normalized)
        errors = verify_normalization(original, normalized, changes, metadata)
        if errors:
            raise ValueError(json.dumps(errors, ensure_ascii=True))
        detail_sheet = workbook["Process Step Detail"]
        headers = [key for key in normalized["Process Step Detail"][0] if key != "_row"]
        for column, header in enumerate(headers, 1):
            detail_sheet.cell(1, column, header)
        for row in normalized["Process Step Detail"]:
            for column, header in enumerate(headers, 1):
                detail_sheet.cell(row["_row"], column, row[header])
        detail_sheet.auto_filter.ref = detail_sheet.dimensions
        for name, rows in (("BPMN Steps", metadata["steps"]), ("BPMN RACI", metadata["raci"]),
                           ("BPMN Control References", metadata["controls"]), ("Validation Issues", metadata["issues"]),
                           ("Change Log", changes)):
            write_sheet(workbook, name, rows)
        issue_counts = dict(sorted(Counter(row["Code"] for row in metadata["issues"]).items()))
        summary = {
            "source": str(source), "source_sha256": source_hash,
            "normalized_workbook": destination.name, "source_preserved": True,
            "structural_validation": "PASS", "normalization_validation": "PASS",
            "excel_round_trip_validation": "PASS",
            "bpmn_execution_readiness": "NEEDS_BUSINESS_CONFIRMATION",
            "bpmn_xml_and_render_validation": "NOT_RUN_NO_DIAGRAMS_GENERATED",
            "source_counts": before["counts"], "changed_cells": len(changes),
            "changes_by_column": dict(Counter(row["Column"] for row in changes)),
            "normalized_steps": len(metadata["steps"]), "raci_variants": len(metadata["raci"]),
            "issue_counts": issue_counts, "structural_errors": [],
            "rules": ["Source sheets retained; edits fully logged; original owner/RACI/detail columns preserved.",
                      "Existing 20 variants remain distinct; no source rows removed.",
                      "People lanes use human teams; Maker/Checker duties stay separate.",
                      "Application names, IDs, qualifiers and existing mapping rows remain unchanged.",
                      "Mutually exclusive TRUE/otherwise tax branches use XOR, not inclusive OR.",
                      "Missing apps/data stay missing; no new hosting assignments or BDEs invented.",
                      "Control references are not a complete BPMN flow graph; display order is not execution order.",
                      "Prior process-specific timeout/parking decisions are not silently approved for every variant.",
                      "Source PDF/hash claims are retained but not independently verified against unavailable documents."],
        }
        write_sheet(workbook, "Validation Summary", [{"Check": key, "Result": value} for key, value in summary.items()])
        output_dir.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_suffix(".tmp.xlsx")
        workbook.save(temporary)
        reopened = openpyxl.load_workbook(temporary, data_only=False)
        try:
            saved = tables_from(reopened)
            saved_source = {name: saved[name] for name in original}
            errors = verify_normalization(original, saved_source, changes, metadata)
            for name, expected in (("BPMN Steps", metadata["steps"]), ("BPMN RACI", metadata["raci"]),
                                   ("BPMN Control References", metadata["controls"]), ("Validation Issues", metadata["issues"]),
                                   ("Change Log", changes)):
                actual = [{key: value for key, value in row.items() if key != "_row"} for row in saved[name]]
                wanted = [{key: excel_value(value) if value != "" else None for key, value in row.items()} for row in expected]
                if actual != wanted:
                    errors.append({"error": "Excel round-trip mismatch", "sheet": name})
            if errors:
                raise ValueError(json.dumps(errors, ensure_ascii=True))
        finally:
            reopened.close()
        if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
            raise ValueError("Source workbook changed during processing")
        temporary.replace(destination)
        (output_dir / "validation.json").write_text(json.dumps(summary, indent=2, ensure_ascii=True), encoding="utf-8")
        document = {"validation": summary, **metadata, "changes": changes,
                    "tables": {name: [{key: value for key, value in row.items() if key != "_row"} for row in rows]
                               for name, rows in normalized.items()}}
        (output_dir / "normalized_process_data.json").write_text(json.dumps(document, indent=2, ensure_ascii=True), encoding="utf-8")
        report = ["# Synthetic Process Data Validation", "", "## Result", "",
                  "Data integrity, normalization and Excel round-trip: **PASS**.",
                  "Executable BPMN readiness: **needs business confirmation**. No BPMN XML or rendering was validated.", "",
                  f"Original: `{source.name}` (unchanged). SHA-256: `{source_hash}`.",
                  f"Corrected workbook: [{destination.name}]({destination.name}).", "",
                  f"{len(original['Process'])} process variants; {len(metadata['steps'])} steps; {len(metadata['raci'])} distinct per-step RACI variants.",
                  f"{len(changes)} source-cell corrections, all recorded in Change Log. Source owner/RACI/detail values retained alongside normalized values.", "",
                  "## Corrections", ""]
        report.extend(f"- {key}: {value} cells." for key, value in summary["changes_by_column"].items())
        report.extend(["", "## Rules", "", *[f"- {rule}" for rule in summary["rules"]], "", "## Remaining Issues", "",
                       "These counts are affected process/step records, not structural errors. Full IDs and explanations are in Validation Issues and normalized_process_data.json.", "",
                       "| Issue | Count |", "|---|---:|"])
        report.extend(f"| {code} | {count} |" for code, count in issue_counts.items())
        report.extend(["", "## Per-Process Coverage", "", "| Process ID | Name | Steps | App gaps | Data gaps |", "|---|---|---:|---:|---:|"])
        for process in original["Process"]:
            records = [row for row in metadata["steps"] if row["Process ID"] == process["Process ID"]]
            report.append(f"| {process['Process ID']} | {process['L4 Process Name']} | {len(records)} | {sum(not row['Application IDs'] for row in records)} | {sum(not row['Data Mapping Count'] for row in records)} |")
        report.extend(["", "## Usage", "", "Use BPMN Steps for human lanes and source-linked application capabilities; use BPMN RACI for all duty-preserving variants.",
                       "Keep Application and Step Application Mapping as inventory/many-to-many links, not sequence flows.",
                       "BPMN Control References preserves textual rules and scoped step references, not inferred executable edges.",
                       "Approve the open decisions, then build and validate BPMN XML/DI separately. The existing single-process renderer is unchanged.", "",
                       "Rerun: `python normalize_process_workbook.py Synthetic_Hackathon_Data_10.xlsx --output-dir synthetic_process_data`.", "Requires openpyxl.", ""])
        (output_dir / "validation.md").write_text("\n".join(report), encoding="utf-8")
        print(json.dumps(summary, indent=2, ensure_ascii=True))
        return summary
    finally:
        workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Audit and normalize synthetic process metadata for BPMN.")
    parser.add_argument("source", nargs="?", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "synthetic_process_data")
    arguments = parser.parse_args()
    run(arguments.source, arguments.output_dir)