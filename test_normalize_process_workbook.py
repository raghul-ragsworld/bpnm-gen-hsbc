from collections import Counter
from contextlib import redirect_stdout
from copy import deepcopy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

from normalize_process_workbook import (
    DEFAULT_SOURCE, audit, build_metadata, normalize, normalize_assignments,
    parse_raci, run, tables_from, verify_normalization,
)


class WorkbookNormalizationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        workbook = openpyxl.load_workbook(DEFAULT_SOURCE)
        try:
            cls.source = tables_from(workbook)
        finally:
            workbook.close()

    def test_source_integrity_and_variant_coverage(self):
        self.assertEqual(audit(self.source)["structural_errors"], [])
        counts = Counter(row["Process ID"] for row in self.source["Process Step"])
        self.assertEqual(len(counts), 20)
        self.assertEqual(set(counts.values()), {29})

    def test_dangling_and_cross_process_references_are_rejected(self):
        for value in ("AOB-PR-9999-V01-S004", "AOB-PR-1001-V02-S004"):
            with self.subTest(value=value):
                tables = deepcopy(self.source)
                tables["Process Step Detail"][0]["Detail Value"] = "route to " + value
                self.assertTrue(audit(tables)["structural_errors"])

    def test_bad_application_mapping_is_rejected(self):
        tables = deepcopy(self.source)
        tables["Step Application Mapping"][0]["Application ID"] = "APP-NOT-FOUND"
        self.assertTrue(audit(tables)["structural_errors"])

    def test_mapping_to_wrong_detail_is_rejected(self):
        tables = deepcopy(self.source)
        tables["Step Application Mapping"][0]["PSD ID"] = tables["Process Step Detail"][1]["PSD ID"]
        self.assertTrue(audit(tables)["structural_errors"])

    def test_duplicate_primary_and_order_keys_are_rejected(self):
        for field in ("Process Step ID", "Reorder"):
            with self.subTest(field=field):
                tables = deepcopy(self.source)
                tables["Process Step"][1][field] = tables["Process Step"][0][field]
                self.assertTrue(audit(tables)["structural_errors"])

    def test_raci_rejects_duplicate_missing_or_empty_roles(self):
        for value in ("R=Maker,R=Checker,A=Lead,C=Ops,I=Requestor", "R=Maker,A=Lead", "R=,A=Lead,C=Ops,I=Requestor"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_raci(value)

    def test_maker_checker_duties_and_unknown_ops(self):
        maker = normalize_assignments("Onboarding Ops (Maker)", "R=Maker,A=Ops Lead,C=Checker,I=Ops")
        checker = normalize_assignments("Process Ops (Checker)", "R=Checker,A=Ops Lead,C=Maker,I=Requestor")
        self.assertEqual(maker[0], "Onboarding Ops")
        self.assertEqual(checker[0], "Process Ops")
        self.assertEqual(maker[2], {"R": "Maker", "C": "Checker"})
        self.assertEqual(checker[2], {"R": "Checker", "C": "Maker"})
        self.assertIn("I=Operations Team (unresolved)", maker[1])

    def test_conflicting_responsible_owner_is_not_silently_replaced(self):
        with self.assertRaises(ValueError):
            normalize_assignments("Onboarding Ops (Maker)", "R=Platform,A=Lead,C=Ops,I=Requestor")

    def test_only_proven_exclusive_gateways_are_changed(self):
        normalized, changes = normalize(self.source)
        corrections = [row for row in changes if row["Column"] == "Detail Type"]
        self.assertEqual(len(corrections), 20)
        self.assertTrue(all(row["Before"] == "Gateway Rule (Inclusive)" and row["After"] == "Gateway Rule (XOR)" for row in corrections))
        tables = deepcopy(self.source)
        gateway = next(row for row in tables["Process Step Detail"] if row["Detail Type"] == "Gateway Rule (Inclusive)")
        gateway["Detail Value"] = "Run either or both applicable branches"
        with self.assertRaises(ValueError):
            normalize(tables)

    def test_missing_values_and_application_alternatives_are_preserved(self):
        normalized, changes = normalize(self.source)
        metadata = build_metadata(normalized)
        self.assertEqual(verify_normalization(self.source, normalized, changes, metadata), [])
        self.assertEqual(len(metadata["steps"]), 580)
        self.assertEqual(len(metadata["raci"]), 620)
        self.assertEqual(sum(not row["Application IDs"] for row in metadata["steps"]), 300)
        self.assertEqual(sum(not row["Data Mapping Count"] for row in metadata["steps"]), 120)
        for name in ("Application", "Step Application Mapping", "Process Step Data", "Step Attribute Mapping", "Source Process"):
            self.assertEqual(normalized[name], self.source[name])
        transfer = next(row for row in metadata["steps"] if row["Process Step ID"] == "AOB-PR-1001-V01-S008")
        self.assertEqual(transfer["Application IDs"], ["APP-003", "APP-004"])

    def test_source_provenance_loss_is_detected(self):
        normalized, changes = normalize(self.source)
        metadata = build_metadata(normalized)
        normalized["Process Step Detail"][0]["Source RACI"] = "R=Other,A=Other,C=Other,I=Other"
        self.assertTrue(verify_normalization(self.source, normalized, changes, metadata))

    def test_saved_deliverables_and_source_hash(self):
        original_hash = hashlib.sha256(DEFAULT_SOURCE.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory() as directory, redirect_stdout(io.StringIO()):
            summary = run(DEFAULT_SOURCE, Path(directory))
            self.assertEqual(summary["changed_cells"], 2340)
            self.assertEqual(summary["excel_round_trip_validation"], "PASS")
            workbook = openpyxl.load_workbook(Path(directory) / summary["normalized_workbook"])
            try:
                checks = {row[0]: row[1] for row in workbook["Validation Summary"].iter_rows(min_row=2, values_only=True)}
                self.assertEqual(checks["excel_round_trip_validation"], "PASS")
                self.assertEqual(workbook["Process Step Detail"].max_row, 1241)
                self.assertEqual(workbook["BPMN RACI"].max_row, 621)
            finally:
                workbook.close()
            saved = json.loads((Path(directory) / "normalized_process_data.json").read_text(encoding="utf-8"))
            self.assertEqual(saved["validation"], summary)
            self.assertEqual(summary["bpmn_xml_and_render_validation"], "NOT_RUN_NO_DIAGRAMS_GENERATED")
        self.assertEqual(hashlib.sha256(DEFAULT_SOURCE.read_bytes()).hexdigest(), original_hash)


if __name__ == "__main__":
    unittest.main()